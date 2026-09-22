"""Blank, stable-key questionnaires that round-trip through reviewed import."""
import io
import math
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4
from reportlab.lib.colors import HexColor
from reportlab.lib.utils import simpleSplit
from .twin import fields


def excel():
    book = Workbook()
    instructions = book.active
    instructions.title = '00_Anleitung'
    for line in ['RAUTE · Profil-Fragebogen', 'Im Blatt Fragebogen die gelben Antwortfelder ausfüllen.',
                 'IDs und Spaltenüberschriften unverändert lassen. Unbeantwortete Fragen bleiben leer.',
                 'Quelle, URL und Notizen nach Möglichkeit ergänzen. Datei als .xlsx speichern.',
                 'Anschließend im Profil hochladen, Vorschau prüfen und gewünschte Antworten übernehmen.',
                 'Die Vorlage enthält keine bestehenden Profilangaben.']:
        instructions.append([line])
    instructions.column_dimensions['A'].width = 110
    sheet = book.create_sheet('Fragebogen')
    headers = ['ID', 'Bereich', 'Frage', 'Anforderung', 'Antwort / Input', 'Quelle / Dokument', 'Source URL', 'Notizen']
    sheet.append(headers)
    for field in fields():
        sheet.append([field['key'], field['section']+' · '+field['group'], field['label'], field['requirement'], '', '', '', ''])
    widths = [18, 36, 75, 18, 85, 40, 45, 50]
    for i, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(i)].width = width
    for row in sheet:
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if cell.row == 1:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='5742CE')
            elif cell.column >= 5:
                cell.fill = PatternFill('solid', fgColor='FFF5D6')
                cell.number_format = '@'
        sheet.row_dimensions[row[0].row].height = 48 if row[0].row > 1 else 30
    sheet.freeze_panes = 'E2'
    sheet.auto_filter.ref = sheet.dimensions
    out = io.BytesIO(); book.save(out)
    return out.getvalue()


def pdf():
    out = io.BytesIO()
    c = canvas.Canvas(out, pagesize=A4)
    c.setTitle('RAUTE Profil-Fragebogen')
    c.setKeywords('raute-questionnaire-v1')
    width, height = A4
    catalog = fields()
    pages = math.ceil(len(catalog)/3)
    for offset in range(0, len(catalog), 3):
        c.setFont('Helvetica-Bold', 15)
        c.drawString(36, height-36, 'RAUTE · Profil-Fragebogen')
        c.setFont('Helvetica', 9)
        c.drawString(36, height-53, 'Digital ausfüllen, mit Formulardaten speichern und im Profil wieder hochladen.')
        c.drawString(36, height-67, 'Unbekannte Antworten leer lassen. Quellen beziehen sich auf die eingetragene Antwort.')
        y = height-92
        for field in catalog[offset:offset+3]:
            c.setFont('Helvetica', 8)
            c.drawString(36, y, f"{field['key']} · {field['requirement']} · {field['group']}"[:115])
            c.setFont('Helvetica-Bold', 10)
            lines = simpleSplit(field['label'], 'Helvetica-Bold', 10, width-72)
            for line in lines:
                y -= 13; c.drawString(36, y, line)
            y -= 14
            for kind, label, box_height, multiline in [('answer','Antwort',52,True),('source','Quelle / Dokument',18,False),('url','Quellen-URL',18,False),('note','Notizen',18,False)]:
                c.setFont('Helvetica', 8);c.drawString(36, y, label);y -= box_height+4
                c.acroForm.textfield(name=f'{kind}__{field["key"]}', tooltip=f'{field["key"]}: {label}',
                    x=36,y=y,width=width-72,height=box_height,fontName='Helvetica',fontSize=10,
                    borderWidth=.5,borderColor=HexColor('#D5D3E3'),fillColor=HexColor('#FAFAFE'),
                    fieldFlags='multiline' if multiline else '',maxlen=20000 if kind=='answer' else 2000)
                y -= 13
            y -= 7
        c.setFont('Helvetica', 8);c.drawRightString(width-36, 18, f'{offset//3+1} / {pages}')
        c.showPage()
    c.save()
    return out.getvalue()
